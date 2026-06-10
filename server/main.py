### backend/main.py
import asyncio
import re
from fastapi import FastAPI, UploadFile, File, HTTPException, Body, Form
from fastapi.middleware.cors import CORSMiddleware
from dotenv import load_dotenv
from typing import List, Dict, Any, Optional, Tuple
from pydantic import BaseModel
import pdfplumber
import io
import os
from services import BasicTransactionAnalyzer
from services import AITransactionAnalyzer
from services import AIMFAnalyzer
from services import AIMStockAnalyzer
from services import AIMWealthDistributionAnalyzer
import logging
from openpyxl import load_workbook
from datetime import datetime, date
from services.stock_holdings_parser import parse_stock_holdings_excel
from services.stock_market_data import (
    build_manual_redistribution_hints,
    enrich_symbols_parallel,
)
from services.cas_portfolio import parse_cas_full_pipeline

# Load environment variables from .env file
load_dotenv()

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class TransactionRequest(BaseModel):
    transactions: List[List[Any]]  # Changed from List[List[str]] to handle null values
    use_ai: bool = False

class MFAnalysisRequest(BaseModel):
    mf_payload: Dict[str, Any]

class MFAskRequest(BaseModel):
    mf_payload: Dict[str, Any]
    question: str
    history: Optional[List[Dict[str, str]]] = None


class StockAnalysisRequest(BaseModel):
    stock_payload: Dict[str, Any]


class StockAskRequest(BaseModel):
    stock_payload: Dict[str, Any]
    question: str
    history: Optional[List[Dict[str, str]]] = None


class StockPortfolioForEnrichment(BaseModel):
    summary: Dict[str, Any]
    holdings: List[Dict[str, Any]]
    computed: Dict[str, Any]


class WealthDistributionRequest(BaseModel):
    wealth_payload: Dict[str, Any]


class WealthDistributionAskRequest(BaseModel):
    wealth_payload: Dict[str, Any]
    question: str
    history: Optional[List[Dict[str, str]]] = None


class NorthStarSnapshotRequest(BaseModel):
    inputs: Dict[str, Any]
    career_profile: Optional[Dict[str, Any]] = None
    horizon_profile: Optional[str] = "balanced"
    target_milestone_inr: Optional[float] = None
    risk_tolerance: Optional[float] = None
    user_id: Optional[str] = "default"


class NorthStarAskRequest(BaseModel):
    north_star_payload: Dict[str, Any]
    question: str
    history: Optional[List[Dict[str, str]]] = None


class WealthComputeRequest(BaseModel):
    inputs: Dict[str, Any]


class StartupOpportunityRequest(BaseModel):
    skill_profile: Dict[str, Any]
    top_n: Optional[int] = 10


class KnowledgeGraphQueryRequest(BaseModel):
    template: str
    params: Optional[Dict[str, Any]] = None


class CopilotSymbolRequest(BaseModel):
    symbol: str
    market: Optional[str] = "IN"
    horizon: Optional[str] = "medium"
    position_context: Optional[Dict[str, Any]] = None


class CopilotPortfolioRequest(BaseModel):
    stock_payload: Dict[str, Any]
    horizon: Optional[str] = "medium"
    refresh_enrichment: Optional[bool] = False


class CopilotAskRequest(BaseModel):
    copilot_payload: Dict[str, Any]
    question: str
    history: Optional[List[Dict[str, str]]] = None


app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def _rows_from_tables(page_tables) -> List[List[Any]]:
    rows: List[List[Any]] = []
    for table in (page_tables or []):
        rows.extend(table)
    return rows


def _extract_tables_from_page(page) -> List[List[Any]]:
    """
    Strategy priority:
    1. "lines" — uses drawn grid lines; produces clean, well-bounded cells.
       Used whenever it yields at least one non-empty row.
    2. "text" — groups text by x/y proximity; catches borderless tables
       (e.g. transaction pages without ruled lines) but may split cells.
    Only fall back to "text" when "lines" finds nothing for the page.
    """
    # Primary: explicit ruled lines (clean cells, correct for display)
    try:
        line_tables = page.extract_tables(
            {"vertical_strategy": "lines", "horizontal_strategy": "lines"}
        )
        line_rows = _rows_from_tables(line_tables)
        # If lines strategy found substantive content, use it
        non_empty = [r for r in line_rows if any(c for c in r if c)]
        if non_empty:
            return line_rows
    except Exception:
        pass

    # Fallback: text-alignment strategy (handles borderless tables)
    try:
        text_tables = page.extract_tables(
            {"vertical_strategy": "text", "horizontal_strategy": "text"}
        )
        return _rows_from_tables(text_tables)
    except Exception:
        pass

    return []


# Column name aliases reused for header detection in the upload layer
_TRANSACTION_HEADER_ALIASES = {
    "date":        ["tran date", "date", "transaction date", "txn date", "value date"],
    "particulars": ["particulars", "description", "narration", "details", "transaction remarks"],
    "amount":      [
        "debit", "withdrawal", "withdrawal amt.", "dr amount", "debit amount",
        "credit", "deposit", "deposit amt.", "cr amount", "credit amount",
    ],
}


def _normalise_cell(s: Any) -> str:
    return " ".join(str(s).strip().split()).lower()


_SWEEP_SUBSTRINGS = ["sweep trf", "sweep transfer", "sweep int", "closure proceeds"]
_FLEXI_INT_RE_MAIN = re.compile(r'\d{8,}\s+int:', re.IGNORECASE)


def _cell_contains_sweep(cell: Any) -> bool:
    return any(kw in str(cell).lower() for kw in _SWEEP_SUBSTRINGS)


def _find_transaction_header_idx(rows: List[List[Any]]) -> int:
    """
    Return the index of the first row that looks like a transaction-table
    header (contains date + particulars aliases). Falls back to 0 if not found.
    """
    for idx, row in enumerate(rows):
        cells = [_normalise_cell(c) for c in row]
        has_date = any(c in _TRANSACTION_HEADER_ALIASES["date"]       for c in cells)
        has_part = any(c in _TRANSACTION_HEADER_ALIASES["particulars"] for c in cells)
        has_amt  = any(c in _TRANSACTION_HEADER_ALIASES["amount"]      for c in cells)
        if has_date and has_part:
            return idx
        if has_date and has_amt:
            return idx
    return 0


def _find_particulars_col(header_row: List[Any]) -> Optional[int]:
    """Return the index of the Particulars-like column in the header row."""
    for i, cell in enumerate(header_row):
        if _normalise_cell(cell) in _TRANSACTION_HEADER_ALIASES["particulars"]:
            return i
    return None


_DATE_RE   = re.compile(r'\d{2}-\d{2}-\d{4}|\d{4}-\d{2}-\d{2}')
_AMOUNT_RE = re.compile(r'^\d{1,3}(?:,\d{2,3})*\.\d{2}$')


def _has_date_or_amount(row: List[Any]) -> bool:
    """True if this row is a 'main' transaction row (carries date or amount)."""
    for cell in row:
        s = str(cell).strip()
        if not s:
            continue
        if _DATE_RE.match(s) or _AMOUNT_RE.match(s):
            return True
    return False


_SUMMARY_LABELS = {"summary", "summary:", "summary :", "summary:"}


def _truncate_at_summary(rows: List[List[Any]]) -> List[List[Any]]:
    """
    Drop everything from the first 'Summary' row onwards.
    Bank statements print a footer block (Total Debits / Total Credits / FD
    scheme tables) after the last transaction row — none of that should be
    displayed or analysed.
    """
    for idx, row in enumerate(rows):
        for cell in row:
            if str(cell).strip().lower() in _SUMMARY_LABELS:
                return rows[:idx]
    return rows


def _filter_sweep_rows(rows: List[List[Any]]) -> List[List[Any]]:
    """
    Remove sweep transactions (and their description rows) from the flat
    table list using a lookahead buffer.

    Row layout in this bank's PDF:
      [blank description rows]  ← precede the main row
      [main row: date + amount]
      [blank tail rows]         ← follow the main row (account number etc.)

    We buffer blank rows and flush them onto the NEXT main row.
    If that main row's combined particulars contain a sweep keyword, we
    discard all its rows (buffer + main row) instead of appending them.
    """
    # Find the particulars column so we know where to look for sweep text
    if not rows:
        return rows
    part_col = _find_particulars_col(rows[0])  # rows[0] is already the header

    filtered: List[List[Any]] = [rows[0]]     # always keep the header
    pending: List[List[Any]] = []             # buffered pre-description rows

    for row in rows[1:]:
        if _has_date_or_amount(row):
            # Collect all particulars text seen so far (from buffer + this row)
            parts_text = " ".join(
                str(r[part_col]).strip()
                for r in pending + [row]
                if part_col is not None and part_col < len(r) and str(r[part_col]).strip()
            ).lower()

            # Normalise whitespace so "Closure\nProceeds" matches "closure proceeds"
            parts_text = " ".join(parts_text.split())
            is_sweep = (
                any(kw in parts_text for kw in _SWEEP_SUBSTRINGS)
                or bool(_FLEXI_INT_RE_MAIN.search(parts_text))
            )
            if is_sweep:
                # Discard this sweep / flexi-interest transaction and its buffered rows
                pending = []
            else:
                filtered.extend(pending)
                filtered.append(row)
                pending = []
        else:
            # Blank / description / tail row — buffer it
            pending.append(row)

    # Flush any trailing rows that follow the last main row
    if pending:
        # Check if leftover pending rows themselves are sweep-related
        raw_parts = " ".join(
            str(r[part_col]).strip()
            for r in pending
            if part_col is not None and part_col < len(r) and str(r[part_col]).strip()
        ).lower()
        parts_text = " ".join(raw_parts.split())  # normalise embedded newlines
        if not any(kw in parts_text for kw in _SWEEP_SUBSTRINGS):
            filtered.extend(pending)

    return filtered


@app.post("/upload-pdf")
async def upload_pdf(file: UploadFile = File(...)):
    raw = await file.read()
    tables: List[List[Any]] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        logger.info("PDF pages: %d", len(pdf.pages))
        for page in pdf.pages:
            tables.extend(_extract_tables_from_page(page))
    logger.info("Total rows extracted from PDF: %d", len(tables))

    # Trim pre-transaction header rows (bank info block)
    start = _find_transaction_header_idx(tables)
    if start > 0:
        logger.info("Trimming %d pre-transaction rows (header at row %d)", start, start)
        tables = tables[start:]

    # Drop the Summary footer (Total Debits / Total Credits rows + FD tables)
    before_summary = len(tables)
    tables = _truncate_at_summary(tables)
    if len(tables) < before_summary:
        logger.info("Summary truncation removed %d footer rows (%d → %d)",
                    before_summary - len(tables), before_summary, len(tables))

    # Remove sweep / flexi-account transfer rows so they never reach the UI
    before = len(tables)
    tables = _filter_sweep_rows(tables)
    logger.info("Sweep filter removed %d rows (%d → %d)", before - len(tables), before, len(tables))

    return {"tables": tables}

def _safe_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    s = s.replace(",", "")
    try:
        return float(s)
    except Exception:
        return None

def _safe_percent(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    if s.endswith("%"):
        s = s[:-1].strip()
    s = s.replace(",", "")
    try:
        return float(s)
    except Exception:
        return None

def _find_first_text_cell(rows: List[List[Any]], text: str) -> Optional[Tuple[int, int]]:
    target = text.strip().lower()
    for r_idx, row in enumerate(rows):
        for c_idx, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower() == target:
                return (r_idx, c_idx)
    return None

def _as_on_from_heading(value: Any) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    # Expected: "HOLDINGS AS ON YYYY-MM-DD"
    prefix = "holdings as on"
    lower = s.lower()
    if prefix in lower:
        idx = lower.find(prefix) + len(prefix)
        rest = s[idx:].strip()
        # Try ISO date first
        try:
            d = datetime.strptime(rest, "%Y-%m-%d").date()
            return d.isoformat()
        except Exception:
            pass
        # Try DD-MM-YYYY
        try:
            d = datetime.strptime(rest, "%d-%m-%Y").date()
            return d.isoformat()
        except Exception:
            return rest
    return None

@app.post("/upload-mf-holdings-excel")
async def upload_mf_holdings_excel(file: UploadFile = File(...)) -> Dict[str, Any]:
    try:
        raw = await file.read()
        wb = load_workbook(filename=io.BytesIO(raw), data_only=True)
        if "Holdings" in wb.sheetnames:
            ws = wb["Holdings"]
        else:
            ws = wb[wb.sheetnames[0]]

        rows: List[List[Any]] = []
        max_col = ws.max_column or 1
        for r in range(1, (ws.max_row or 1) + 1):
            row = []
            for c in range(1, max_col + 1):
                row.append(ws.cell(row=r, column=c).value)
            rows.append(row)

        # Profile
        profile: Dict[str, Any] = {}
        for label, key in [("Name", "name"), ("Mobile Number", "mobile"), ("PAN", "pan")]:
            pos = _find_first_text_cell(rows, label)
            if pos:
                r, c = pos
                val = rows[r][c + 1] if c + 1 < len(rows[r]) else None
                if val is not None and str(val).strip():
                    profile[key] = str(val).strip()

        # Summary
        summary: Dict[str, Any] = {}
        as_on: Optional[str] = None
        for row in rows:
            if as_on:
                break
            for cell in row:
                maybe = _as_on_from_heading(cell)
                if maybe:
                    as_on = maybe
                    break
        if as_on:
            summary["as_on"] = as_on

        summary_header_pos = None
        for i, row in enumerate(rows):
            if any(isinstance(cell, str) and cell.strip().lower() == "total investments" for cell in row):
                summary_header_pos = i
                break
        if summary_header_pos is not None and summary_header_pos + 1 < len(rows):
            header_row = [str(x).strip() if x is not None else "" for x in rows[summary_header_pos]]
            value_row = rows[summary_header_pos + 1]
            for idx, h in enumerate(header_row):
                v = value_row[idx] if idx < len(value_row) else None
                h_lower = h.lower()
                if h_lower == "total investments":
                    summary["total_invested"] = _safe_float(v)
                elif h_lower == "current portfolio value":
                    summary["current_value"] = _safe_float(v)
                elif h_lower == "profit/loss":
                    summary["profit_loss"] = _safe_float(v)
                elif h_lower == "profit/loss %":
                    summary["profit_loss_pct"] = _safe_percent(v)
                elif h_lower == "xirr":
                    summary["xirr_pct"] = _safe_percent(v)

        # Holdings table
        holdings: List[Dict[str, Any]] = []
        holdings_header_idx = None
        for i, row in enumerate(rows):
            if any(isinstance(cell, str) and cell.strip().lower() == "scheme name" for cell in row):
                holdings_header_idx = i
                break
        if holdings_header_idx is None:
            raise HTTPException(status_code=400, detail="Could not find holdings header row (Scheme Name).")

        header = [str(x).strip() if x is not None else "" for x in rows[holdings_header_idx]]
        col = {h.lower(): idx for idx, h in enumerate(header) if h}

        def get(row: List[Any], key: str) -> Any:
            idx = col.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        for row in rows[holdings_header_idx + 1 :]:
            scheme = get(row, "scheme name")
            if scheme is None or (isinstance(scheme, str) and not scheme.strip()):
                continue
            entry = {
                "scheme_name": str(scheme).strip(),
                "amc": (str(get(row, "amc")).strip() if get(row, "amc") is not None else None),
                "category": (str(get(row, "category")).strip() if get(row, "category") is not None else None),
                "sub_category": (str(get(row, "sub-category")).strip() if get(row, "sub-category") is not None else None),
                "folio_no": (str(get(row, "folio no.")).strip() if get(row, "folio no.") is not None else None),
                "source": (str(get(row, "source")).strip() if get(row, "source") is not None else None),
                "units": _safe_float(get(row, "units")),
                "invested_value": _safe_float(get(row, "invested value")),
                "current_value": _safe_float(get(row, "current value")),
                "returns_value": _safe_float(get(row, "returns")),
                "xirr_pct": _safe_percent(get(row, "xirr")),
            }
            holdings.append(entry)

        # Computed allocations & concentration
        def sum_by(key: str) -> List[Dict[str, Any]]:
            buckets: Dict[str, float] = {}
            for h in holdings:
                k = h.get(key) or "Unknown"
                v = h.get("current_value") or 0.0
                buckets[str(k)] = buckets.get(str(k), 0.0) + float(v)
            out = [{"name": k, "value": round(v, 2)} for k, v in buckets.items()]
            out.sort(key=lambda x: x["value"], reverse=True)
            return out

        total_current = sum((h.get("current_value") or 0.0) for h in holdings) or 0.0
        top_schemes: Dict[str, float] = {}
        for h in holdings:
            name = h.get("scheme_name") or "Unknown"
            top_schemes[str(name)] = top_schemes.get(str(name), 0.0) + float(h.get("current_value") or 0.0)
        top_schemes_list = [{"name": k, "value": round(v, 2), "pct": round((v / total_current) * 100, 2) if total_current else 0.0} for k, v in top_schemes.items()]
        top_schemes_list.sort(key=lambda x: x["value"], reverse=True)

        computed = {
            "counts": {
                "rows": len(holdings),
                "unique_schemes": len(set(h.get("scheme_name") for h in holdings)),
                "unique_amcs": len(set(h.get("amc") for h in holdings)),
                "unique_folios": len(set(h.get("folio_no") for h in holdings)),
                "unique_sources": len(set(h.get("source") for h in holdings)),
            },
            "total_current_from_rows": round(float(total_current), 2),
            "allocation_by_category": sum_by("category"),
            "allocation_by_sub_category": sum_by("sub_category"),
            "allocation_by_amc": sum_by("amc"),
            "allocation_by_source": sum_by("source"),
            "top_schemes_by_current_value": top_schemes_list[:10],
        }

        return {"profile": profile, "summary": summary, "holdings": holdings, "computed": computed}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error parsing MF holdings excel: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/upload-stock-holdings-excel")
async def upload_stock_holdings_excel(file: UploadFile = File(...)) -> Dict[str, Any]:
    """
    Parse Excel only (fast). Call POST /enrich-stock-portfolio next for Yahoo data.
    """
    try:
        raw = await file.read()
        return parse_stock_holdings_excel(raw)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.error(f"Error parsing stock holdings excel: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/parse-cas-pdf")
async def parse_cas_pdf(
    file: UploadFile = File(...),
    password: str = Form(""),
    prefer_commercial: str = Form("false"),
    include_raw: str = Form("false"),
) -> Dict[str, Any]:
    """
    Parse India CAS PDF (CAMS, KFintech, CDSL eCAS, NSDL eCAS) using open-source casparser
    on the server. Optional CAS_PARSER_API_KEY enables cloud fallback via api.casparser.in
    (free tier available) for difficult files or extra asset-class coverage.
    Password is usually your PAN (CDSL/NSDL may use encrypted-PAN format as per statement).
    """
    try:
        raw = await file.read()
        if len(raw) > 25 * 1024 * 1024:
            raise HTTPException(
                status_code=400,
                detail="PDF too large (max 25 MB).",
            )
        pc = prefer_commercial.strip().lower() in ("1", "true", "yes", "on")
        ir = include_raw.strip().lower() in ("1", "true", "yes", "on")
        return await asyncio.to_thread(
            parse_cas_full_pipeline,
            raw,
            password or "",
            prefer_commercial=pc,
            include_raw=ir,
        )
    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except RuntimeError as e:
        raise HTTPException(status_code=503, detail=str(e))
    except Exception as e:
        logger.error("CAS parse failed: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/enrich-stock-portfolio")
async def enrich_stock_portfolio(body: StockPortfolioForEnrichment) -> Dict[str, Any]:
    try:
        parsed = body.model_dump()
        symbols = sorted(
            {
                str(h.get("symbol"))
                for h in parsed.get("holdings") or []
                if h.get("symbol")
            }
        )
        if len(symbols) > 80:
            raise HTTPException(
                status_code=400,
                detail="Too many symbols (max 80). Split the file or trim rows.",
            )
        batches = max(1, (len(symbols) + 3) // 4)
        overall_timeout_sec = float(min(600, max(90, 50 * batches)))
        enrichment = enrich_symbols_parallel(
            symbols,
            max_workers=4,
            per_future_timeout_sec=45,
            overall_timeout_sec=overall_timeout_sec,
        )
        manual_hints = build_manual_redistribution_hints(
            parsed.get("computed", {}).get("allocation_by_symbol") or [],
            enrichment,
        )
        return {"enrichment": enrichment, "manual_hints": manual_hints}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error enriching stock portfolio: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/analyze-basic-transactions")
async def analyze_basic_transactions(request: TransactionRequest) -> Dict[str, Any]:
    try:
        # Convert None/null values to empty strings and ensure all values are strings
        processed_transactions = [
            [str(cell) if cell is not None else "" for cell in row]
            for row in request.transactions
        ]

        basic_analyzer = BasicTransactionAnalyzer()
        basic_analysis = basic_analyzer.analyze_transactions(processed_transactions)

        if "error" in basic_analysis:
            logger.error("analyze_transactions error: %s | all rows: %s",
                         basic_analysis["error"], processed_transactions)
            raise HTTPException(status_code=422, detail=basic_analysis["error"])

        logger.info(
            "analyze_transactions: %d input rows → %d categorised txns "
            "(sweep filtered: %d) | total_debits=%s total_credits=%s",
            len(processed_transactions),
            len(basic_analysis.get("categorized_transactions", [])),
            basic_analysis.get("sweep_filtered_count", 0),
            basic_analysis.get("total_debits"),
            basic_analysis.get("total_credits"),
        )

        return {"basic_analysis": basic_analysis}
    except HTTPException:
        raise  # let 4xx pass through unchanged
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# Separate API endpoint for AITransactionAnalyzer
@app.post("/analyze-ai-transactions")
async def analyze_ai_transactions(request: TransactionRequest) -> Dict[str, Any]:
    api_key = os.getenv('OPENAI_API_KEY')
    if not api_key:
        logger.error("OPENAI_API_KEY not found in environment variables")
        raise ValueError("OPENAI_API_KEY not found in environment variables")

    try:
        logger.info("Received request for AI transaction analysis")
        logger.info(f"Request data: {request.transactions}")

        # Convert None/null values to empty strings and ensure all values are strings
        processed_transactions = [
            [str(cell) if cell is not None else "" for cell in row]
            for row in request.transactions
        ]

        logger.info("Processed transactions for AI analysis")

        ai_analyzer = AITransactionAnalyzer(api_key)
        ai_analysis = await ai_analyzer.analyze_transactions(processed_transactions)

        logger.info("AI analysis completed successfully")
        return {"ai_analysis": ai_analysis}
    except Exception as e:
        logger.error(f"Error during AI transaction analysis: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/analyze-mf-ai")
async def analyze_mf_ai(request: MFAnalysisRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        logger.error("OPENAI_API_KEY not found in environment variables")
        raise ValueError("OPENAI_API_KEY not found in environment variables")

    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        mf_analyzer = AIMFAnalyzer(api_key=api_key, model=model)
        report = await mf_analyzer.generate_report(request.mf_payload)
        return {"report": report}
    except Exception as e:
        logger.error(f"Error during MF AI analysis: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ask-mf-ai")
async def ask_mf_ai(request: MFAskRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        logger.error("OPENAI_API_KEY not found in environment variables")
        raise ValueError("OPENAI_API_KEY not found in environment variables")

    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        mf_analyzer = AIMFAnalyzer(api_key=api_key, model=model)
        answer = await mf_analyzer.ask(
            mf_payload=request.mf_payload,
            question=request.question,
            history=request.history,
        )
        return answer
    except Exception as e:
        logger.error(f"Error during MF Ask AI: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/analyze-stock-ai")
async def analyze_stock_ai(request: StockAnalysisRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        logger.error("OPENAI_API_KEY not found in environment variables")
        raise ValueError("OPENAI_API_KEY not found in environment variables")

    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        stock_analyzer = AIMStockAnalyzer(api_key=api_key, model=model)
        report = await stock_analyzer.generate_report(request.stock_payload)
        return {"report": report}
    except Exception as e:
        logger.error(f"Error during stock AI analysis: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ask-stock-ai")
async def ask_stock_ai(request: StockAskRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        logger.error("OPENAI_API_KEY not found in environment variables")
        raise ValueError("OPENAI_API_KEY not found in environment variables")

    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        stock_analyzer = AIMStockAnalyzer(api_key=api_key, model=model)
        answer = await stock_analyzer.ask(
            stock_payload=request.stock_payload,
            question=request.question,
            history=request.history,
        )
        return answer
    except Exception as e:
        logger.error(f"Error during stock Ask AI: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/analyze-wealth-distribution-ai")
async def analyze_wealth_distribution_ai(request: WealthDistributionRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        logger.error("OPENAI_API_KEY not found in environment variables")
        raise ValueError("OPENAI_API_KEY not found in environment variables")

    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        wd = AIMWealthDistributionAnalyzer(api_key=api_key, model=model)
        report = await wd.generate_report(request.wealth_payload)
        return {"report": report}
    except Exception as e:
        logger.error(f"Error during wealth distribution AI: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ask-wealth-distribution-ai")
async def ask_wealth_distribution_ai(request: WealthDistributionAskRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        logger.error("OPENAI_API_KEY not found in environment variables")
        raise ValueError("OPENAI_API_KEY not found in environment variables")

    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        wd = AIMWealthDistributionAnalyzer(api_key=api_key, model=model)
        answer = await wd.ask(
            payload=request.wealth_payload,
            question=request.question,
            history=request.history,
        )
        return answer
    except Exception as e:
        logger.error(f"Error during wealth distribution Ask AI: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


# --- North Star Mode ---

@app.post("/wealth/compute")
async def wealth_compute(request: WealthComputeRequest) -> Dict[str, Any]:
    try:
        from services import wealth_engine
        return wealth_engine.compute(request.inputs)
    except Exception as e:
        logger.error(f"Error in wealth compute: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/wealth/simulate")
async def wealth_simulate(request: WealthComputeRequest) -> Dict[str, Any]:
    try:
        from services import wealth_engine, wealth_simulation_engine
        metrics = wealth_engine.compute(request.inputs)
        return wealth_simulation_engine.run_all(request.inputs, metrics)
    except Exception as e:
        logger.error(f"Error in wealth simulate: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/startup/opportunities")
async def startup_opportunities(request: StartupOpportunityRequest) -> Dict[str, Any]:
    try:
        from services import startup_opportunity_engine
        return startup_opportunity_engine.run(request.skill_profile, request.top_n or 10)
    except Exception as e:
        logger.error(f"Error in startup opportunities: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/north-star/recommend")
async def north_star_recommend(request: NorthStarSnapshotRequest) -> Dict[str, Any]:
    try:
        from services import north_star_engine, recommendation_engine, wealth_engine, wealth_simulation_engine, startup_opportunity_engine
        from services.north_star_engine import _portfolio_from_inputs, _market_conditions
        metrics = wealth_engine.compute(request.inputs)
        simulation = wealth_simulation_engine.run_all(request.inputs, metrics)
        career = request.career_profile or {}
        startups = startup_opportunity_engine.run({
            "skills": career.get("skills", []),
            "wealth_context": request.inputs,
            "runway_months": metrics.get("emergency_fund_months"),
            "risk_tolerance": request.risk_tolerance or 50,
        })
        ctx = {
            "inputs": request.inputs,
            "metrics": metrics,
            "simulation": simulation,
            "startup_opportunities": startups,
            "career_profile": career,
            "portfolio": _portfolio_from_inputs(request.inputs),
            "market_conditions": _market_conditions(),
        }
        return recommendation_engine.recommend(ctx, request.horizon_profile or "balanced")
    except Exception as e:
        logger.error(f"Error in north star recommend: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/compute-north-star")
async def compute_north_star(request: NorthStarSnapshotRequest) -> Dict[str, Any]:
    try:
        from services import north_star_engine
        snapshot = request.model_dump()
        return north_star_engine.compute(snapshot)
    except Exception as e:
        logger.error(f"Error in compute north star: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/knowledge-graph/query")
async def knowledge_graph_query(request: KnowledgeGraphQueryRequest) -> Dict[str, Any]:
    try:
        from services import knowledge_graph
        return knowledge_graph.query_graph(request.template, request.params)
    except Exception as e:
        logger.error(f"Error in knowledge graph query: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/knowledge-graph/snapshot")
async def knowledge_graph_snapshot() -> Dict[str, Any]:
    try:
        from services import knowledge_graph
        return knowledge_graph.get_snapshot()
    except Exception as e:
        logger.error(f"Error in knowledge graph snapshot: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/analyze-north-star-ai")
async def analyze_north_star_ai(request: NorthStarAskRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=400, detail="OPENAI_API_KEY not configured")
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        from services.analyze_north_star_ai import AINorthStarAnalyzer
        analyzer = AINorthStarAnalyzer(api_key=api_key, model=model)
        report = await analyzer.generate_report(request.north_star_payload)
        return {"report": report}
    except Exception as e:
        logger.error(f"Error during north star AI: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ask-north-star-ai")
async def ask_north_star_ai(request: NorthStarAskRequest) -> Dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=400, detail="OPENAI_API_KEY not configured")
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        from services.analyze_north_star_ai import AINorthStarAnalyzer
        analyzer = AINorthStarAnalyzer(api_key=api_key, model=model)
        answer = await analyzer.ask(request.north_star_payload, request.question, request.history)
        return answer
    except Exception as e:
        logger.error(f"Error during north star Ask AI: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


# --- Investment Copilot ---

@app.post("/copilot/analyze-symbol")
async def copilot_analyze_symbol(request: CopilotSymbolRequest) -> Dict[str, Any]:
    try:
        from services import investment_copilot
        return investment_copilot.analyze_symbol(
            request.symbol,
            request.market or "IN",
            request.horizon or "medium",  # type: ignore
            request.position_context,
        )
    except Exception as e:
        logger.error(f"Error in copilot analyze symbol: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/copilot/analyze-portfolio")
async def copilot_analyze_portfolio(request: CopilotPortfolioRequest) -> Dict[str, Any]:
    try:
        from services import investment_copilot
        return investment_copilot.analyze_portfolio(
            request.stock_payload,
            request.horizon or "medium",  # type: ignore
            request.refresh_enrichment or False,
        )
    except Exception as e:
        logger.error(f"Error in copilot analyze portfolio: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))