"""
Parse equity holdings from broker-exported Excel (Zerodha-like and generic layouts).
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook


def _safe_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    s = s.replace(",", "")
    if s.endswith("%"):
        s = s[:-1].strip()
    try:
        return float(s)
    except Exception:
        return None


def normalize_equity_symbol(raw: str) -> str:
    s = str(raw).strip().upper()
    if not s:
        return ""
    for suf in ("-EQ", "-BE", "-BL", "-BZ", "-SM", "-ST"):
        if s.endswith(suf):
            s = s[: -len(suf)]
            break
    s = re.sub(r"\s+", "", s)
    return s


def _row_strings(row: List[Any]) -> List[str]:
    out: List[str] = []
    for cell in row:
        if cell is None:
            out.append("")
        else:
            out.append(str(cell).strip())
    return out


def _normalize_header_cell(cell: str) -> str:
    s = cell.strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _cell_matches_synonym(cell_norm: str, alts: List[str]) -> bool:
    if not cell_norm:
        return False
    if cell_norm in alts:
        return True
    for a in alts:
        if len(a) <= 3:
            continue
        if cell_norm.startswith(a):
            return True
        if a.endswith("*") and cell_norm.startswith(a[:-1]):
            return True
    return False


def _build_col_map_for_row(lowered: List[str]) -> Dict[str, int]:
    synonyms = {
        "symbol": [
            "symbol",
            "scrip",
            "stock",
            "instrument",
            "ticker",
            "company symbol",
            "stock symbol",
        ],
        "name": [
            "name",
            "company name",
            "stock name",
            "security name",
            "instrument name",
        ],
        "isin": ["isin"],
        "qty": [
            "qty",
            "quantity",
            "qty.",
            "net qty",
            "total qty",
            "balance",
            "shares",
            "quantity.",
        ],
        "avg_price": [
            "avg. price",
            "avg price",
            "average price",
            "avg buy price",
            "avg. buy price",
            "average buy price",
            "avg. cost",
            "average cost",
        ],
        "invested": [
            "invested",
            "invested value",
            "buy value",
            "total invested",
            "amount invested",
            "investment",
            "cost",
            "total cost",
        ],
        "current_value": [
            "current value",
            "cur. value",
            "cur value",
            "curr. value",
            "curr value",
            "market value",
            "current val",
            "mkt val",
            "mkt. val",
            "valuation",
            "value (inr)",
            "present value",
            "total value",
        ],
        "ltp": [
            "ltp",
            "last price",
            "last traded price",
            "cmp",
            "close price",
            "closing price",
            "last close",
        ],
        "pnl_pct": [
            "p&l %",
            "p/l %",
            "return %",
            "returns %",
            "ret %",
            "gain %",
            "profit/loss %",
            "pnl %",
        ],
        "pnl_abs": [
            "p&l",
            "p/l",
            "profit/loss",
            "gain",
            "returns",
        ],
    }

    col_map: Dict[str, int] = {}
    for c_idx, raw in enumerate(lowered):
        cell = _normalize_header_cell(raw)
        if not cell:
            continue
        for key, alts in synonyms.items():
            if key in col_map:
                continue
            if _cell_matches_synonym(cell, alts):
                col_map[key] = c_idx
                break
    return col_map


def _header_row_score(col_map: Dict[str, int]) -> int:
    if "symbol" not in col_map:
        return -1
    score = len(col_map) * 2
    if "current_value" in col_map:
        score += 8
    if "ltp" in col_map:
        score += 5
    if "qty" in col_map:
        score += 3
    if "invested" in col_map:
        score += 2
    if "pnl_pct" in col_map:
        score += 1
    return score


def _find_best_holdings_header(rows: List[List[Any]]) -> Optional[Tuple[int, Dict[str, int]]]:
    best: Optional[Tuple[int, Dict[str, int], int]] = None
    for r_idx, row in enumerate(rows):
        lowered = [_normalize_header_cell(x) for x in _row_strings(row)]
        if not any(lowered):
            continue
        col_map = _build_col_map_for_row(lowered)
        if "symbol" not in col_map:
            continue
        if not ("current_value" in col_map or "ltp" in col_map or "qty" in col_map):
            continue
        sc = _header_row_score(col_map)
        if best is None or sc > best[2]:
            best = (r_idx, col_map, sc)
    if best:
        return best[0], best[1]
    return None


def _derive_current_value(
    qty: Optional[float],
    invested: Optional[float],
    pnl_pct: Optional[float],
    ltp: Optional[float],
    raw_current: Optional[float],
) -> Optional[float]:
    if raw_current is not None:
        return raw_current
    if qty is not None and ltp is not None and qty > 0 and ltp > 0:
        return round(qty * ltp, 2)
    if invested is not None and pnl_pct is not None and invested > 0:
        return round(invested * (1.0 + pnl_pct / 100.0), 2)
    return None


def parse_stock_holdings_excel(raw: bytes) -> Dict[str, Any]:
    wb = load_workbook(filename=io.BytesIO(raw), data_only=True)
    if "Holdings" in wb.sheetnames:
        ws = wb["Holdings"]
    else:
        ws = wb[wb.sheetnames[0]]

    rows: List[List[Any]] = []
    max_col = ws.max_column or 1
    for r in range(1, (ws.max_row or 1) + 1):
        row = []
        for c in range(1, max_col + 1):
            row.append(ws.cell(row=r, column=c).value)
        rows.append(row)

    found = _find_best_holdings_header(rows)
    if not found:
        raise ValueError(
            "Could not find a holdings table. Expected columns like Symbol, Qty, "
            "Current value or LTP (or Invested + P&L %)."
        )

    header_idx, col = found

    def get(row: List[Any], key: str) -> Any:
        idx = col.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    holdings: List[Dict[str, Any]] = []
    derived_value_note = False
    for row in rows[header_idx + 1 :]:
        sym_raw = get(row, "symbol")
        sym = normalize_equity_symbol(str(sym_raw)) if sym_raw else ""
        if not sym or sym in ("TOTAL", "SUBTOTAL", "GRAND TOTAL"):
            continue
        qty = _safe_float(get(row, "qty"))
        if qty is not None and qty == 0:
            continue

        ltp = _safe_float(get(row, "ltp"))
        raw_cv = _safe_float(get(row, "current_value"))
        invested = _safe_float(get(row, "invested"))
        pnl_pct = _safe_float(get(row, "pnl_pct"))

        current_value = _derive_current_value(
            qty, invested, pnl_pct, ltp, raw_cv
        )
        if current_value is not None and raw_cv is None:
            derived_value_note = True

        entry = {
            "symbol": sym,
            "name": (
                str(get(row, "name")).strip()
                if get(row, "name") is not None
                else None
            ),
            "isin": (
                str(get(row, "isin")).strip()
                if get(row, "isin") is not None
                else None
            ),
            "qty": qty,
            "avg_price": _safe_float(get(row, "avg_price")),
            "invested_value": invested,
            "current_value": current_value,
            "pnl_pct": pnl_pct,
            "pnl_abs": _safe_float(get(row, "pnl_abs")),
        }
        if raw_cv is None and current_value is not None:
            if (
                ltp is not None
                and qty is not None
                and qty > 0
                and abs(current_value - qty * ltp) < 1.0
            ):
                entry["value_source"] = "ltp_x_qty"
            elif (
                invested is not None
                and pnl_pct is not None
                and invested > 0
            ):
                entry["value_source"] = "invested_and_pnl_pct"
        holdings.append(entry)

    total_current = sum(float(h.get("current_value") or 0) for h in holdings)
    total_invested = sum(float(h.get("invested_value") or 0) for h in holdings)
    by_symbol: Dict[str, float] = {}
    for h in holdings:
        s = h.get("symbol") or "?"
        by_symbol[s] = by_symbol.get(s, 0.0) + float(h.get("current_value") or 0)

    alloc = [
        {
            "name": k,
            "value": round(v, 2),
            "pct": round((v / total_current) * 100, 2) if total_current else 0.0,
        }
        for k, v in by_symbol.items()
    ]
    alloc.sort(key=lambda x: x["value"], reverse=True)

    weights = [a["pct"] / 100.0 for a in alloc]
    hhi = sum(w * w for w in weights) if weights else 0.0
    top3 = round(sum(a["pct"] for a in alloc[:3]), 2) if alloc else 0.0

    summary = {
        "total_holdings_rows": len(holdings),
        "unique_symbols": len(by_symbol),
        "total_current_value": round(total_current, 2),
        "total_invested_value": round(total_invested, 2),
        "parse_note": (
            "Some current values were estimated from LTP×Qty or Invested×(1+P&L%) "
            "because those cells were empty (common with Excel formulas — use Paste Values)."
            if derived_value_note
            else None
        ),
    }

    computed = {
        "allocation_by_symbol": alloc,
        "concentration": {
            "hhi": round(hhi, 4),
            "top_3_weight_pct": round(top3, 2),
            "largest_position_pct": round(alloc[0]["pct"], 2) if alloc else 0.0,
        },
    }

    return {
        "summary": summary,
        "holdings": holdings,
        "computed": computed,
    }
